"""250문항 샘플 Excel을 현재 PostgreSQL 문제 테이블에 안전하게 적재한다."""

import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from models import DifficultyLevel, Question, QuestionType, SessionLocal, Subject, engine


EXCEL_PATH = Path(__file__).with_name("문제은행_샘플데이터_250문항_1.xlsx")
SOURCE_FILE = EXCEL_PATH.name

SUBJECT_NAMES = {
    "JAVA": "Java",
    "JS": "JavaScript",
    "SQL": "SQL",
    "SPRING": "Spring",
}

REQUIRED_HEADERS = [
    "id", "subject_code", "subject_unit", "content", "code_snippet",
    "difficulty", "benchmark_time", "option_1", "option_2", "option_3",
    "option_4", "correct_option", "explanation", "tags",
]
REQUIRED_HEADERS_WITH_REQUIRED = REQUIRED_HEADERS + ["is_required"]


def ensure_columns() -> None:
    statements = [
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS external_id INTEGER",
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS subject_code VARCHAR(50)",
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS subject_unit VARCHAR(255)",
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS code_snippet TEXT",
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS benchmark_time INTEGER",
        "ALTER TABLE questions ADD COLUMN IF NOT EXISTS correct_option INTEGER",
        "CREATE INDEX IF NOT EXISTS idx_questions_external_id ON questions(external_id)",
        "CREATE INDEX IF NOT EXISTS idx_questions_subject_code ON questions(subject_code)",
    ]
    with engine.begin() as connection:
        for statement in statements:
            connection.execute(text(statement))


def import_questions() -> tuple[int, int]:
    workbook = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    if headers not in (REQUIRED_HEADERS, REQUIRED_HEADERS_WITH_REQUIRED):
        raise ValueError(f"Excel headers do not match: {headers}")

    db = SessionLocal()
    inserted = 0
    skipped = 0
    try:
        subjects = {}
        for code, name in SUBJECT_NAMES.items():
            subject = db.query(Subject).filter(Subject.name == name).first()
            if subject is None:
                subject = Subject(name=name, description=f"{name} 샘플 문제")
                db.add(subject)
                db.flush()
            subjects[code] = subject

        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            exists = db.query(Question.id).filter(
                Question.source_file == SOURCE_FILE,
                Question.external_id == int(row["id"]),
            ).first()
            if exists:
                skipped += 1
                continue

            code = str(row["subject_code"]).upper()
            options = [row[f"option_{number}"] for number in range(1, 5)]
            correct_option = int(row["correct_option"])
            question = Question(
                subject_id=subjects[code].id,
                question_type=QuestionType.EXCEL,
                title=str(row["subject_unit"]),
                content=str(row["content"]),
                correct_answer=str(correct_option - 1),
                explanation=str(row["explanation"]) if row["explanation"] is not None else None,
                difficulty=DifficultyLevel(str(row["difficulty"]).upper()),
                source_file=SOURCE_FILE,
                options=json.dumps(options, ensure_ascii=False),
                tags=str(row["tags"]),
                external_id=int(row["id"]),
                subject_code=code,
                subject_unit=str(row["subject_unit"]),
                code_snippet=str(row["code_snippet"]) if row["code_snippet"] is not None else None,
                benchmark_time=int(row["benchmark_time"]),
                correct_option=correct_option,
                is_required=str(row.get("is_required", "false")).strip().lower() in {"true", "1", "yes", "y", "필수"},
            )
            db.add(question)
            inserted += 1

        db.commit()
        return inserted, skipped
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    ensure_columns()
    created, existing = import_questions()
    print(f"inserted={created} skipped={existing}")
